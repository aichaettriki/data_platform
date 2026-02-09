"""
Spark Job for Competitive Scores Processing with MinIO Integration
Reads from MinIO bucket, processes data, and saves back to MinIO
"""
 
import pandas as pd
import logging
from openpyxl import load_workbook
from pyspark.sql import SparkSession
from pyspark.sql import functions as F
from pyspark.sql.window import Window
import sys
import os
import re
from io import BytesIO
from datetime import datetime
 
# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)
 
 
def create_spark_session(app_name="Competitif Scores Processor"):
    """Create Spark session with MinIO configuration"""
    spark = (
        SparkSession.builder
        .appName(app_name)
        .config("spark.hadoop.fs.s3a.access.key", os.getenv("MINIO_ROOT_USER", "minioadmin"))
        .config("spark.hadoop.fs.s3a.secret.key", os.getenv("MINIO_ROOT_PASSWORD", "minioadmin"))
        .config("spark.hadoop.fs.s3a.endpoint", os.getenv("MINIO_ENDPOINT", "http://minio:9000"))
        .config("spark.hadoop.fs.s3a.path.style.access", "true")
        .config("spark.hadoop.fs.s3a.impl", "org.apache.hadoop.fs.s3a.S3AFileSystem")
        .config("spark.hadoop.fs.s3a.connection.ssl.enabled", "false")
        .config("spark.sql.sources.partitionOverwriteMode", "dynamic")
        .config("spark.driver.memory", "4g")
        .config("spark.executor.memory", "4g")
        .getOrCreate()
    )
    return spark
 
 
def find_latest_files_in_minio(minio_client, bucket, base_folder):
    """
    Find latest Excel files in the specified MinIO folder
    Similar to your find_latest_files_recursive but for MinIO
    """
    logger.info(f"Searching for Excel files in: {bucket}/{base_folder}")
   
    objects = minio_client.list_objects(bucket, prefix=base_folder, recursive=True)
   
    excel_files = []
    for obj in objects:
        if obj.object_name.endswith(('.xlsx', '.xls')) and not obj.object_name.split('/')[-1].startswith('~'):
            excel_files.append(obj.object_name)
            logger.info(f"  Found: {obj.object_name}")
   
    if not excel_files:
        raise FileNotFoundError(f"No Excel files found in {bucket}/{base_folder}")
   
    # Sort by name (latest first) - you can customize this logic
    excel_files.sort(reverse=True)
   
    logger.info(f"\n✅ Total Excel files found: {len(excel_files)}")
    logger.info(f"   Using: {excel_files[0]}")
   
    return excel_files[0]
 

def clean_roman_prefix(text: str) -> str:
    """
    Supprime un préfixe en chiffre romain au début d'un titre.
    Exemple:
    'III - Innovation' -> 'Innovation'
    'IV. Emploi' -> 'Emploi'
    """
    roman_pattern = r'^\s*(?P<roman>M{0,4}(CM|CD|D?C{0,3})'
    roman_pattern += r'(XC|XL|L?X{0,3})(IX|IV|V?I{0,3}))'
    roman_pattern += r'[\s\-\.\:\/]+'

    return re.sub(roman_pattern, '', text, flags=re.IGNORECASE).strip()

 
class MinIOConfig:
    """MinIO connection configuration from environment variables"""
    def __init__(self):
        self.endpoint = os.getenv("MINIO_ENDPOINT", "http://minio:9000")
        self.access_key = os.getenv("MINIO_ROOT_USER", "minioadmin")
        self.secret_key = os.getenv("MINIO_ROOT_PASSWORD", "minioadmin")
        self.bucket_raw = "01-raw"
        self.bucket_transformed = "02-transformed"
 
 
class CompetitifScoresProcessor:
    """Main processor for competitive scores data with MinIO integration"""
   
    def __init__(self, year, month, day):
        self.year = year
        self.month = month
        self.day = day
       
        # MinIO configuration
        self.minio_config = MinIOConfig()
       
        # Build MinIO paths
        self.input_path = f"{year}/{month}/ITCEQ/competitivite/positionnement/"
        self.output_path = "ITCEQ/Compétivité/Positionnement/"
       
        # Detection configuration
        self.detection_config = {
            'column_b_must_be_empty': True,
            'excluded_texts': ['Score', 'Rang', 'Pays', 'Rank', 'Country', 'Year'],
            'min_text_length': 3,
        }
       
        # Initialize Spark with MinIO support
        self.spark = create_spark_session("Competitif Scores Processor")
        logger.info("Spark session initialized with MinIO support")
       
        # Initialize MinIO client (for non-Spark operations)
        self._init_minio_client()
   
    def _init_minio_client(self):
        """Initialize MinIO client for direct object operations"""
        from minio import Minio
       
        # Remove http:// or https:// from endpoint
        endpoint = self.minio_config.endpoint.replace("http://", "").replace("https://", "")
       
        self.minio_client = Minio(
            endpoint,
            access_key=self.minio_config.access_key,
            secret_key=self.minio_config.secret_key,
            secure=False  # Set to True if using HTTPS
        )
        logger.info(f"MinIO client initialized: {endpoint}")
   
    def find_excel_file(self):
        """Find the Excel file in MinIO bucket using discovery logic"""
        logger.info("="*80)
        logger.info("SEARCHING FOR EXCEL FILE IN MINIO")
        logger.info("="*80)
        logger.info(f"Bucket: {self.minio_config.bucket_raw}")
        logger.info(f"Path: {self.input_path}")
       
        # Use the file discovery function
        self.input_file_key = find_latest_files_in_minio(
            self.minio_client,
            self.minio_config.bucket_raw,
            self.input_path
        )
       
        logger.info(f"\n✅ Using file: {self.input_file_key}")
       
        return self.input_file_key
   
    def extract_data_from_minio(self):
        """Extract data from Excel file in MinIO"""
        logger.info("="*80)
        logger.info("STEP 1: DATA EXTRACTION FROM MINIO")
        logger.info("="*80)
        logger.info(f"Reading: s3a://{self.minio_config.bucket_raw}/{self.input_file_key}")
       
        # Download file from MinIO to memory
        response = self.minio_client.get_object(
            self.minio_config.bucket_raw,
            self.input_file_key
        )
       
        # Read Excel file from bytes
        excel_data = BytesIO(response.read())
        response.close()
        response.release_conn()
       
        # Determine sheet name
        wb = load_workbook(excel_data, read_only=True)
        sheet_names = wb.sheetnames
        logger.info(f"Available sheets: {sheet_names}")
       
        # Look for the competitive positioning sheet
        sheet_name = None
        for name in sheet_names:
            if 'positionnement' in name.lower() and 'compétitif' in name.lower():
                sheet_name = name
                break
       
        if not sheet_name:
            sheet_name = sheet_names[0]  # Default to first sheet
       
        logger.info(f"Using sheet: {sheet_name}")
        self.sheet_name = sheet_name
       
        # Reset BytesIO position for pandas
        excel_data.seek(0)
       
        # Read with pandas
        df = pd.read_excel(excel_data, sheet_name=sheet_name, header=None)
        logger.info(f"Data shape: {df.shape}")
       
        # Store for later use
        self.excel_data_bytes = excel_data
        self.raw_data = df
       
        return df
   
    def identify_titles(self):
        """
        Ultra-flexible title/indicator detection
        """
        logger.info("="*80)
        logger.info("STEP 2: TITLE/INDICATOR DETECTION")
        logger.info("="*80)
       
        # Reset BytesIO for openpyxl
        self.excel_data_bytes.seek(0)
       
        wb = load_workbook(self.excel_data_bytes)
        ws = wb[self.sheet_name]
       
        detected_titles = []
       
        for row_idx in range(1, ws.max_row + 1):
            cell_a = ws.cell(row=row_idx, column=1)
            cell_b = ws.cell(row=row_idx, column=2)
           
            if not cell_a.value or not isinstance(cell_a.value, str):
                continue
           
            text = str(cell_a.value).strip()
            text = clean_roman_prefix(text)
            if len(text) < self.detection_config['min_text_length']:
                continue
           
            if cell_b.value is not None:
                continue
           
            if text in self.detection_config['excluded_texts']:
                continue
           
            font = cell_a.font
            is_bold = font.bold if font else False
            font_size = font.size if font else 10.0
           
            detected_titles.append({
                'row_index': row_idx - 1,
                'title': text,
                'font_size': font_size,
                'is_bold': is_bold,
            })
           
            bold_marker = "BOLD" if is_bold else "    "
            logger.info(f"✅ Row {row_idx:3d} | {font_size:4.1f}pt | {bold_marker} | {text}")
       
        logger.info("="*80)
        logger.info(f"TOTAL TITLES/INDICATORS DETECTED: {len(detected_titles)}")
        logger.info("="*80)
       
        if len(detected_titles) == 0:
            raise ValueError("No titles detected!")
       
        self.title_sections = detected_titles
        return detected_titles
   
    def transform_data(self):
        """Transform data from wide to long format"""
        logger.info("="*80)
        logger.info("STEP 3: DATA TRANSFORMATION")
        logger.info("="*80)
    
        df = self.raw_data
        title_sections = self.title_sections
        
        # On définit la date de chargement maintenant pour tout le batch
        loading_date = datetime.now()
    
        all_records = []
    
        for i, section in enumerate(title_sections):
            title = section['title']
            start_row = section['row_index']
        
            if i < len(title_sections) - 1:
                end_row = title_sections[i + 1]['row_index']
            else:
                end_row = len(df)
        
            logger.info(f"\n[{i+1}/{len(title_sections)}] Processing: {title}")
            logger.info(f"    Rows {start_row} to {end_row}")
        
            section_df = df.iloc[start_row:end_row, :].copy()
            section_df.reset_index(drop=True, inplace=True)
        
            try:
                # Find years row
                years_row_idx = None
            
                for row_idx in range(min(10, len(section_df))):
                    row = section_df.iloc[row_idx]
                    year_candidates = []
                
                    for col_idx in range(1, min(6, len(row))):
                        val = row[col_idx]
                        if pd.notna(val):
                            try:
                                num = float(val)
                                if 2000 <= num <= 2030 and num == int(num):
                                    year_candidates.append(int(num))
                            except:
                                pass
                
                    if len(year_candidates) >= 3:
                        if max(year_candidates) - min(year_candidates) <= 20:
                            years_row_idx = row_idx
                            break
            
                if years_row_idx is None:
                    logger.warning(f"    ⚠️  No years row found")
                    continue
            
                # Extract years
                years_row = section_df.iloc[years_row_idx, 1:15]
                years = []
                for y in years_row:
                    if pd.notna(y):
                        try:
                            year_val = int(float(y))
                            if 2000 <= year_val <= 2030:
                                years.append(year_val)
                        except:
                            pass
            
                if not years:
                    logger.warning(f"    ⚠️  No valid years")
                    continue
            
                logger.info(f"    Years: {years[0]}-{years[-1]} ({len(years)} years)")
            
                # Extract country data
                countries_data = section_df.iloc[years_row_idx + 1:, :].copy()
                count = 0
            
                for _, row in countries_data.iterrows():
                    country = row[0]
                
                    if pd.isna(country) or country == '':
                        continue
                
                    if isinstance(country, str):
                        country_str = str(country).strip()
                        if country_str in ['Score', 'Rang', 'Pays', 'Rank', 'Country', 'Year'] or len(country_str) < 2:
                            continue
                
                    scores = row[1:15].tolist()
                
                    for year, score in zip(years, scores):
                        if pd.notna(score) and score != '':
                            try:
                                score_val = float(score)
                                if 0 <= score_val <= 100:
                                    all_records.append({
                                        'pays': str(country).strip(),
                                        'annee': int(year),
                                        'indicateur': title.strip(),
                                        'valeur': score_val,
                                        'date_chargement': loading_date  # <-- Ajout de la colonne ici
                                    })
                                    count += 1
                            except:
                                continue
            
                logger.info(f"    ✅ Extracted {count} records")
        
            except Exception as e:
                logger.error(f"    ❌ Error: {e}")
                import traceback
                logger.error(traceback.format_exc())
                continue
    
        df_transformed = pd.DataFrame(all_records)
    
        if len(df_transformed) == 0:
            raise ValueError("No data transformed!")
    
        before = len(df_transformed)
        # Mise à jour du subset pour inclure la date_chargement dans la vérification des doublons si nécessaire
        # ou rester sur les colonnes métier uniquement :
        df_transformed = df_transformed.drop_duplicates(subset=['pays', 'annee', 'indicateur'])
        after = len(df_transformed)
    
        if before > after:
            logger.warning(f"⚠️  Removed {before - after} duplicate records")
    
        df_transformed = df_transformed.sort_values(['pays', 'annee', 'indicateur'])
    
        logger.info("\n" + "="*80)
        logger.info("TRANSFORMATION SUMMARY")
        logger.info("="*80)
        logger.info(f"Total records: {len(df_transformed):,}")
        logger.info(f"Unique countries: {df_transformed['pays'].nunique()}")
        logger.info(f"Unique indicators: {df_transformed['indicateur'].nunique()}")
        logger.info(f"Year range: {df_transformed['annee'].min()}-{df_transformed['annee'].max()}")
        logger.info(f"Load Date: {loading_date}") # Info log supplémentaire
    
        self.transformed_data = df_transformed
        return df_transformed


    def calculate_rankings(self):
        """Calculate rankings using Spark"""
        logger.info("="*80)
        logger.info("STEP 4: RANKING CALCULATION (SPARK)")
        logger.info("="*80)
       
        spark_df = self.spark.createDataFrame(self.transformed_data)
        logger.info(f"Total records to rank: {spark_df.count():,}")
       
        window_spec = Window.partitionBy("indicateur", "annee").orderBy(F.col("valeur").desc())
        spark_df_ranked = spark_df.withColumn("rang", F.rank().over(window_spec))
       
        df_ranked = spark_df_ranked.toPandas()
       
        logger.info("\nRanking statistics by indicator:")
        for indicator in sorted(df_ranked['indicateur'].unique()):
            df_ind = df_ranked[df_ranked['indicateur'] == indicator]
            avg_rank = df_ind['rang'].mean()
            max_rank = df_ind['rang'].max()
            min_rank = df_ind['rang'].min()
           
            logger.info(f"\n  📊 {indicator}:")
            logger.info(f"     - Rank range: {min_rank} to {max_rank}")
            logger.info(f"     - Average rank: {avg_rank:.1f}")
            logger.info(f"     - Total records: {len(df_ind)}")
       
        df_ranked = df_ranked.sort_values(['pays', 'annee', 'indicateur'])
       
        logger.info("\n" + "="*80)
        logger.info("RANKING CALCULATION COMPLETED")
        logger.info("="*80)
       
        self.ranked_data = df_ranked
        return df_ranked
   
    def save_to_minio(self):
        """Save ranked data to MinIO in Parquet format (stable overwrite)"""
        logger.info("="*80)
        logger.info("STEP 5: SAVING TO MINIO")
        logger.info("="*80)
    
        # Convert to Spark DataFrame
        spark_df_ranked = self.spark.createDataFrame(self.ranked_data)

        # Build output path
        output_s3_path = f"s3a://{self.minio_config.bucket_transformed}/{self.output_path}"
    
        logger.info(f"Output path: {output_s3_path}")
        logger.info(f"Format: Parquet")
        logger.info(f"Records: {spark_df_ranked.count():,}")
        logger.info(f"Columns: {spark_df_ranked.columns}")

        # =========================================================
        # CONFIGURATION S3A COMPATIBLE MINIO (SPARK VANILLA)
        # =========================================================
        self.spark.conf.set("spark.hadoop.mapreduce.fileoutputcommitter.algorithm.version", "2")
        self.spark.conf.set("spark.hadoop.fs.s3a.committer.name", "directory")

        # =========================================================
        # WRITE PARQUET (OVERWRITE)
        # =========================================================
        spark_df_ranked.write \
            .mode("overwrite") \
            .option("fs.s3a.committer.name", "directory") \
            .parquet(output_s3_path)
    
        logger.info("✅ Data saved to MinIO with overwrite enabled")
        logger.info(f"   Bucket: {self.minio_config.bucket_transformed}")
        logger.info(f"   Path: {self.output_path}")
    
        return {
            'bucket': self.minio_config.bucket_transformed,
            'path': self.output_path,
            'records': len(self.ranked_data),
            'format': 'parquet'
        }


   
    def validate_minio_output(self):
        """Validate the Parquet files in MinIO"""
        logger.info("="*80)
        logger.info("STEP 6: VALIDATION")
        logger.info("="*80)
       
        # Read back from MinIO to validate
        output_s3_path = f"s3a://{self.minio_config.bucket_transformed}/{self.output_path}"
       
        try:
            df_validation = self.spark.read.parquet(output_s3_path)
            count = df_validation.count()
            columns = df_validation.columns
           
            logger.info(f"✅ Validation successful!")
            logger.info(f"   Records read back: {count:,}")
            logger.info(f"   Columns: {columns}")
           
            # Check for required columns
            required_columns = ['pays', 'annee', 'indicateur', 'valeur', 'rang']
            missing_columns = [col for col in required_columns if col not in columns]
           
            if missing_columns:
                logger.error(f"❌ Missing columns: {missing_columns}")
                return False
           
            # Check for nulls
            null_counts = df_validation.select([F.count(F.when(F.col(c).isNull(), c)).alias(c) for c in columns]).collect()[0]
            has_nulls = any(null_counts[c] > 0 for c in columns)
           
            if has_nulls:
                logger.warning("⚠️  Found null values:")
                for col in columns:
                    if null_counts[col] > 0:
                        logger.warning(f"   {col}: {null_counts[col]} nulls")
           
            # Show sample data
            logger.info("\nSample data (first 5 rows):")
            df_validation.show(5, truncate=False)
           
            logger.info("\n✅ ALL VALIDATION CHECKS PASSED!")
            return True
           
        except Exception as e:
            logger.error(f"❌ Validation failed: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return False
   
    def run(self):
        """Execute the full pipeline"""
        logger.info("\n" + "🚀"*40)
        logger.info("STARTING MINIO-BASED COMPETITIVE SCORES PIPELINE")
        logger.info("🚀"*40 + "\n")
       
        try:
            # Step 1: Find and extract from MinIO
            self.find_excel_file()
            self.extract_data_from_minio()
           
            # Step 2: Identify titles
            self.identify_titles()
           
            # Step 3: Transform
            self.transform_data()
           
            # Step 4: Calculate rankings
            self.calculate_rankings()
           
            # Step 5: Save to MinIO
            save_results = self.save_to_minio()
           
            # Step 6: Validate
            validation_passed = self.validate_minio_output()
           
            logger.info("\n" + "✅"*40)
            logger.info("PIPELINE COMPLETED SUCCESSFULLY")
            logger.info("✅"*40)
           
            return {
                'status': 'success',
                'validation_passed': validation_passed,
                'input_bucket': self.minio_config.bucket_raw,
                'input_path': self.input_file_key,
                'output_bucket': save_results['bucket'],
                'output_path': save_results['path'],
                'records_processed': save_results['records'],
                'output_format': save_results['format']
            }
           
        except Exception as e:
            logger.error(f"\n❌ PIPELINE FAILED: {e}")
            import traceback
            logger.error(traceback.format_exc())
            raise
       
        finally:
            self.spark.stop()
            logger.info("Spark session stopped")
 
 
def main():
    """Main entry point for Spark job"""
    if len(sys.argv) < 4:
        print("Usage: spark-submit spark_competitif_processor.py <year> <month> <day>")
        print("Example: spark-submit spark_competitif_processor.py 2024 01 15")
        print("\nMinIO configuration should be set via environment variables:")
        print("  - MINIO_ENDPOINT (default: http://minio:9000)")
        print("  - MINIO_ROOT_USER (default: minioadmin)")
        print("  - MINIO_ROOT_PASSWORD (default: minioadmin)")
        sys.exit(1)
   
    year = sys.argv[1]
    month = sys.argv[2]
    day = sys.argv[3]
   
    logger.info(f"Processing date: {year}/{month}/{day}")
    logger.info(f"MinIO Endpoint: {os.getenv('MINIO_ENDPOINT', 'http://minio:9000')}")
    logger.info(f"MinIO User: {os.getenv('MINIO_ROOT_USER', 'minioadmin')}")
   
    # Run processor
    processor = CompetitifScoresProcessor(year, month, day)
    results = processor.run()
   
    print(f"\n{'='*80}")
    print("FINAL RESULTS:")
    print(f"{'='*80}")
    for key, value in results.items():
        print(f"{key}: {value}")
   
    return results
 
 
if __name__ == "__main__":
    main()